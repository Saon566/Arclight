"""Workbook bridge for the Arclight Streamlit investor portal.

Works locally on Windows with LibreOffice installed and on Streamlit Community
Cloud/Linux when packages.txt installs libreoffice-calc.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Mapping, Any

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
SOURCE_FILE = BASE_DIR / "Vyomaa_P3_P4_Capital_Raise_Model_v8(3).xlsx"


def find_soffice() -> str | None:
    """Return the available LibreOffice executable on Windows or Linux."""
    candidates = [
        shutil.which("libreoffice"),
        shutil.which("soffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    return None


def load_workbook_values(path: str | Path = SOURCE_FILE):
    """Load cached formula results for dashboard display."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(
            f"Workbook not found: {path.name}. Put it in the same folder as app.py."
        )
    return load_workbook(path, data_only=True, read_only=False)


def recalculate_workbook(
    scenario: str = "Base case",
    overrides: Mapping[str, Any] | None = None,
) -> Path:
    """Apply input-cell overrides to a copy, recalculate with LibreOffice, return output path."""
    source = Path(SOURCE_FILE)
    if not source.exists():
        raise FileNotFoundError(
            f"Required Excel workbook '{source.name}' is missing from the app directory."
        )

    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(
            "LibreOffice was not found. On Streamlit Community Cloud, add a packages.txt "
            "file containing 'libreoffice-calc', then reboot the app. Locally, install LibreOffice."
        )

    safe_name = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in scenario)
    output_dir = BASE_DIR / "recalculated"
    output_dir.mkdir(parents=True, exist_ok=True)
    final_path = output_dir / f"Arclight_{safe_name}.xlsx"

    # Each run gets its own LibreOffice profile to avoid concurrent profile locks.
    with tempfile.TemporaryDirectory(prefix="arclight_model_") as tmp:
        tmpdir = Path(tmp)
        input_dir = tmpdir / "input"
        converted_dir = tmpdir / "converted"
        profile_dir = tmpdir / "lo_profile"
        input_dir.mkdir()
        converted_dir.mkdir()
        profile_dir.mkdir()
        staged = input_dir / source.name

        wb = load_workbook(source, data_only=False)
        try:
            for address, value in (overrides or {}).items():
                # Addresses use Sheet!Cell format, e.g. Inputs!C16.
                if "!" in address:
                    sheet_name, cell = address.split("!", 1)
                else:
                    # Backward compatibility for the existing app's C16-style keys.
                    sheet_name, cell = "Inputs", address
                if sheet_name not in wb.sheetnames:
                    raise KeyError(f"Worksheet '{sheet_name}' not found in workbook.")
                wb[sheet_name][cell] = value
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = "auto"
            wb.save(staged)
        finally:
            wb.close()

        command = [
            soffice,
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(converted_dir),
            str(staged),
        ]
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
                env={**os.environ, "HOME": str(tmpdir)},
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("LibreOffice recalculation timed out after 180 seconds.") from exc

        converted = converted_dir / source.name
        if result.returncode != 0 or not converted.exists():
            details = (result.stderr or result.stdout or "No details returned by LibreOffice").strip()
            raise RuntimeError(f"LibreOffice could not recalculate the workbook: {details[-1200:]}")

        shutil.copy2(converted, final_path)

    # Validate that the converted workbook can be opened and contains required sheets.
    check = load_workbook(final_path, data_only=True, read_only=True)
    try:
        required = {"Dashboard", "Inputs", "SPV3", "SPV4", "Checks"}
        missing = required.difference(check.sheetnames)
        if missing:
            raise RuntimeError("Recalculated workbook is missing sheets: " + ", ".join(sorted(missing)))
    finally:
        check.close()
    return final_path
