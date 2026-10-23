from pathlib import Path
import shutil
import subprocess
import tempfile
import openpyxl

BASE_DIR = Path(__file__).resolve().parent
MODEL_FILE = BASE_DIR / "Vyomaa_P3_P4_Capital_Raise_Model_v8(3).xlsx"
SOFFICE_CANDIDATES = [
    Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
    Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
]


def find_soffice():
    """Find LibreOffice on Windows or Linux hosting (e.g. Streamlit Community Cloud)."""
    for candidate in SOFFICE_CANDIDATES:
        if candidate.exists():
            return str(candidate)

    # Streamlit Community Cloud runs on Linux; packages.txt installs this executable.
    return shutil.which("libreoffice") or shutil.which("soffice")


def recalculate_workbook(scenario="Base case", overrides=None):
    """
    Recalculate a COPY of the source workbook. Never modifies the original.
    Overrides are Excel input-cell values, e.g. {"C16": 4.04}.
    """
    if not MODEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel model not found: {MODEL_FILE}\n"
            "Place the workbook in the same folder as app.py."
        )

    soffice = find_soffice()
    if not soffice:
        raise FileNotFoundError(
            "LibreOffice was not found. On Streamlit Community Cloud, ensure "
            "packages.txt is in the same directory as app.py and contains "
            "'libreoffice-calc', then reboot/redeploy the app. Locally on Windows, "
            "install LibreOffice or update SOFFICE_CANDIDATES."
        )

    safe_scenario = "".join(
        ch if ch.isalnum() or ch in ("-", "_") else "_"
        for ch in scenario
    )
    output_dir = BASE_DIR / "recalculated"
    output_dir.mkdir(exist_ok=True)
    final_file = output_dir / f"model_{safe_scenario}.xlsx"

    with tempfile.TemporaryDirectory(prefix="arclight_model_") as temp:
        temp_dir = Path(temp)
        input_dir = temp_dir / "input"
        converted_dir = temp_dir / "converted"
        profile_dir = temp_dir / "lo_profile"
        input_dir.mkdir()
        converted_dir.mkdir()

        scenario_input = input_dir / MODEL_FILE.name
        shutil.copy2(MODEL_FILE, scenario_input)

        if overrides:
            wb = openpyxl.load_workbook(scenario_input, data_only=False)
            if "Inputs" not in wb.sheetnames:
                wb.close()
                raise ValueError("The workbook has no 'Inputs' worksheet.")
            ws = wb["Inputs"]
            for cell, value in overrides.items():
                ws[cell] = value
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = "auto"
            wb.save(scenario_input)
            wb.close()

        command = [
            str(soffice),
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--headless",
            "--convert-to", "xlsx",
            "--outdir", str(converted_dir),
            str(scenario_input),
        ]
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=240,
        )
        converted_file = converted_dir / MODEL_FILE.name

        if result.returncode != 0 or not converted_file.exists():
            raise RuntimeError(
                "LibreOffice did not produce the recalculated workbook.\n"
                f"Exit code: {result.returncode}\n"
                f"STDOUT: {result.stdout}\nSTDERR: {result.stderr}"
            )

        shutil.copy2(converted_file, final_file)

    # Basic validation: can open as a workbook and expected sheets exist.
    check_wb = openpyxl.load_workbook(final_file, data_only=True, read_only=True)
    required = {"Dashboard", "Inputs", "SPV3", "SPV4", "HoldCo", "Checks"}
    missing = required.difference(check_wb.sheetnames)
    if missing:
        check_wb.close()
        raise ValueError(f"Missing expected worksheets: {sorted(missing)}")
    check_wb.close()

    return final_file


def load_workbook_values(path):
    return openpyxl.load_workbook(path, data_only=True, read_only=True)
